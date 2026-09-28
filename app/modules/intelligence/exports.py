"""Board-pack exports for the transformation intelligence briefing.

Composes the executive briefing, project performance, benefits, KPIs and
recommendations into a downloadable **XLSX** (openpyxl) or **PDF** (fpdf2). Pure
rendering over the intelligence engines — no schema change, read-only.
"""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING, Any

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

if TYPE_CHECKING:
    from app.modules.intelligence.service import IntelligenceService

_NAVY = "1F4E79"
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _gather(service: IntelligenceService) -> dict[str, Any]:
    """Collect the briefing and supporting rollups from the engines."""
    return {
        "brief": service.transformation.transformation_intelligence(as_of=None),
        "perf": service.performance.transformation_performance(as_of=None),
        "ben": service.variance.benefits_variance_transformation(),
        "kpi": service.variance.kpi_variance_transformation(),
        "recs": service.recommendations.transformation_recommendations(as_of=None),
    }


def _ascii(value: Any) -> str:
    """Coerce to a core-font-safe string (fpdf2 core fonts are latin-1)."""
    return str(value).encode("latin-1", "replace").decode("latin-1")


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------
def _sheet_header(ws: Worksheet, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left")


def board_pack_xlsx(service: IntelligenceService) -> bytes:
    """Render the board pack as an XLSX workbook and return the bytes."""
    d = _gather(service)
    brief, perf, ben, kpi, recs = d["brief"], d["perf"], d["ben"], d["kpi"], d["recs"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Executive Summary"
    ws["A1"] = "ETIP — Transformation Board Pack"
    ws["A1"].font = Font(bold=True, size=16, color=_NAVY)
    summary = [
        ("As of", str(brief.as_of)),
        ("Overall health", brief.health.rag),
        ("Success outlook", f"{brief.success_score}/100 ({brief.success_label})"),
        ("SPI", brief.spi),
        ("CPI", brief.cpi),
        ("Budget (BAC)", str(brief.bac)),
        ("Forecast cost", str(brief.forecast_cost)),
        ("Projects (red / total)", f"{brief.projects_red} / {brief.projects_total}"),
        ("Open issues", brief.open_issues),
        ("Narrative", brief.narrative),
    ]
    for i, (label, value) in enumerate(summary, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90

    ws = wb.create_sheet("Projects")
    _sheet_header(ws, ["Code", "Health", "SPI", "CPI"])
    for r, p in enumerate(perf.breakdown, start=2):
        ws.cell(row=r, column=1, value=p.code)
        ws.cell(row=r, column=2, value=p.rag)
        ws.cell(row=r, column=3, value=p.spi)
        ws.cell(row=r, column=4, value=p.cpi)

    ws = wb.create_sheet("Benefits")
    _sheet_header(ws, ["Metric", "Value"])
    for r, (label, value) in enumerate(
        [
            ("Benefit count", ben.benefit_count),
            ("Total target", str(ben.total_target)),
            ("Total realized", str(ben.total_realized)),
            ("Variance", str(ben.variance)),
            ("Realization %", ben.realization_percent),
            ("Favourable", ben.favourable),
        ],
        start=2,
    ):
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=value)

    ws = wb.create_sheet("KPIs")
    _sheet_header(ws, ["KPI", "Current", "Target", "Attainment %", "On target"])
    for r, k in enumerate(kpi.items, start=2):
        ws.cell(row=r, column=1, value=k.name)
        ws.cell(row=r, column=2, value=str(k.current_value))
        ws.cell(row=r, column=3, value=str(k.target_value))
        ws.cell(row=r, column=4, value=k.attainment_percent)
        ws.cell(row=r, column=5, value=k.on_target)

    ws = wb.create_sheet("Recommendations")
    _sheet_header(ws, ["Priority", "Category", "Title", "Rationale", "Action"])
    for r, rec in enumerate(recs.recommendations, start=2):
        ws.cell(row=r, column=1, value=rec.priority)
        ws.cell(row=r, column=2, value=rec.category)
        ws.cell(row=r, column=3, value=rec.title)
        ws.cell(row=r, column=4, value=rec.rationale)
        ws.cell(row=r, column=5, value=rec.recommended_action)
    for width, col in ((14, "A"), (14, "B"), (40, "C"), (60, "D"), (60, "E")):
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def board_pack_pdf(service: IntelligenceService) -> bytes:
    """Render the board pack as a PDF and return the bytes."""
    d = _gather(service)
    brief, perf, recs = d["brief"], d["perf"], d["recs"]
    ben, kpi = d["ben"], d["kpi"]

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(31, 78, 121)
    pdf.cell(0, 10, "ETIP - Transformation Board Pack", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(
        0, 6, _ascii(f"As of {brief.as_of}  |  CONFIDENTIAL"), new_x=XPos.LMARGIN, new_y=YPos.NEXT
    )
    pdf.ln(3)

    def section(title: str) -> None:
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(31, 78, 121)
        pdf.cell(0, 8, _ascii(title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(30, 30, 30)

    section("Executive summary")
    for label, value in [
        ("Overall health", brief.health.rag),
        ("Success outlook", f"{brief.success_score}/100 ({brief.success_label})"),
        ("SPI / CPI", f"{brief.spi} / {brief.cpi}"),
        ("Budget (BAC) / Forecast", f"{brief.bac} / {brief.forecast_cost}"),
        ("Projects red / total", f"{brief.projects_red} / {brief.projects_total}"),
    ]:
        pdf.cell(55, 6, _ascii(label), border=0)
        pdf.cell(0, 6, _ascii(value), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(1)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(pdf.epw, 5, _ascii(brief.narrative))
    pdf.ln(3)

    section("Projects")
    pdf.set_font("Helvetica", "B", 10)
    for w, t in ((30, "Code"), (40, "Health"), (30, "SPI"), (30, "CPI")):
        pdf.cell(w, 6, t, border=1)
    pdf.ln()
    pdf.set_font("Helvetica", "", 10)
    for p in perf.breakdown:
        pdf.cell(30, 6, _ascii(p.code), border=1)
        pdf.cell(40, 6, _ascii(p.rag), border=1)
        pdf.cell(30, 6, _ascii(p.spi), border=1)
        pdf.cell(30, 6, _ascii(p.cpi), border=1)
        pdf.ln()
    pdf.ln(2)

    section("Benefits & KPIs")
    pdf.cell(
        0,
        6,
        _ascii(
            f"Benefits: realized {ben.total_realized} of target {ben.total_target} "
            f"({ben.realization_percent}%).  KPIs on target: "
            f"{kpi.kpis_on_target}/{kpi.kpi_count} (avg {kpi.average_attainment}%)."
        ),
        new_x=XPos.LMARGIN,
        new_y=YPos.NEXT,
    )
    pdf.ln(2)

    section("Top recommendations")
    for rec in recs.recommendations[:6]:
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pdf.epw, 5, _ascii(f"[{rec.priority.upper()}] {rec.title} ({rec.category})"))
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(90, 90, 90)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pdf.epw, 5, _ascii(f"{rec.rationale} {rec.recommended_action}"))
        pdf.set_text_color(30, 30, 30)
        pdf.ln(1)

    return bytes(pdf.output())
