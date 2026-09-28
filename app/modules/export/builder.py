"""Board-pack export — a formatted multi-sheet Excel snapshot of the
Intelligence Layer, scoped org-wide or to a single portfolio.

This is a point-in-time report (literal values, no formulas), assembled from the
same intelligence the dashboard shows.
"""

from __future__ import annotations

import io
import uuid
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.intelligence.service import IntelligenceService

_NAVY = "0E1B36"
_HEAD = Font(name="Arial", bold=True, color="FFFFFF", size=11)
_TITLE = Font(name="Arial", bold=True, size=16, color=_NAVY)
_LABEL = Font(name="Arial", bold=True, size=10)
_BODY = Font(name="Arial", size=10)
_HEAD_FILL = PatternFill("solid", fgColor=_NAVY)
_RAG = {
    "green": PatternFill("solid", fgColor="1EA97C"),
    "amber": PatternFill("solid", fgColor="E8A33D"),
    "red": PatternFill("solid", fgColor="D64550"),
}


def _headers(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = _HEAD
        c.fill = _HEAD_FILL
        c.alignment = Alignment(horizontal="left", vertical="center")


def _autosize(ws: Worksheet) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 60)


def build_board_pack(service: IntelligenceService, *, portfolio_id: uuid.UUID | None) -> bytes:
    """Return an .xlsx board pack as bytes."""
    perf = (
        service.portfolio_performance(portfolio_id, as_of=None)
        if portfolio_id
        else service.transformation_performance(as_of=None)
    )
    benefits = (
        service.benefits_variance_portfolio(portfolio_id)
        if portfolio_id
        else service.benefits_variance_transformation()
    )
    kpis = (
        service.kpi_variance_portfolio(portfolio_id)
        if portfolio_id
        else service.kpi_variance_transformation()
    )
    recs = (
        service.portfolio_recommendations(portfolio_id, as_of=None)
        if portfolio_id
        else service.transformation_recommendations(as_of=None)
    )
    scope_label = perf.scope_label

    wb = Workbook()

    # --- Summary sheet ---
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "ETIP — Executive Board Pack"
    ws["A1"].font = _TITLE
    ws["A2"] = (
        f"Scope: {scope_label}    ·    As of: {date.today().isoformat()}    ·    CONFIDENTIAL"
    )
    ws["A2"].font = Font(name="Arial", italic=True, size=10, color="5B6B85")
    rows = [
        ("Overall health", perf.health.rag),
        ("Schedule performance index (SPI)", perf.evm.spi),
        ("Cost performance index (CPI)", perf.evm.cpi),
        ("Budget at completion (BAC)", perf.evm.bac),
        ("Estimate at completion (EAC)", perf.evm.eac),
        ("Projects", perf.project_count),
        ("Benefits target", str(benefits.total_target)),
        ("Benefits realized", str(benefits.total_realized)),
        ("Benefits realization %", benefits.realization_percent),
        ("KPIs on target", f"{kpis.kpis_on_target} / {kpis.kpi_count}"),
    ]
    r = 4
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = _LABEL
        vc = ws.cell(row=r, column=2, value=value)
        vc.font = _BODY
        if label == "Overall health" and isinstance(value, str) and value in _RAG:
            vc.fill = _RAG[value]
            vc.font = Font(name="Arial", bold=True, color="FFFFFF")
        r += 1
    _autosize(ws)

    # --- Projects sheet ---
    ws = wb.create_sheet("Projects")
    _headers(ws, 1, ["Code", "Health", "SPI", "CPI"])
    for i, b in enumerate(perf.breakdown, start=2):
        ws.cell(row=i, column=1, value=b.code).font = _BODY
        hc = ws.cell(row=i, column=2, value=b.rag)
        hc.font = _BODY
        if b.rag in _RAG:
            hc.fill = _RAG[b.rag]
            hc.font = Font(name="Arial", bold=True, color="FFFFFF")
        ws.cell(row=i, column=3, value=b.spi).font = _BODY
        ws.cell(row=i, column=4, value=b.cpi).font = _BODY
    _autosize(ws)

    # --- KPIs sheet ---
    ws = wb.create_sheet("KPIs")
    _headers(ws, 1, ["KPI", "Unit", "Current", "Target", "Attainment %", "On target"])
    for i, k in enumerate(kpis.items, start=2):
        ws.cell(row=i, column=1, value=k.name).font = _BODY
        ws.cell(row=i, column=2, value=k.unit).font = _BODY
        ws.cell(row=i, column=3, value=str(k.current_value)).font = _BODY
        ws.cell(row=i, column=4, value=str(k.target_value)).font = _BODY
        ws.cell(row=i, column=5, value=k.attainment_percent).font = _BODY
        ws.cell(row=i, column=6, value="Yes" if k.on_target else "No").font = _BODY
    _autosize(ws)

    # --- Recommendations sheet ---
    ws = wb.create_sheet("Recommendations")
    _headers(ws, 1, ["Priority", "Category", "Title", "Rationale", "Recommended action"])
    for i, rec in enumerate(recs.recommendations, start=2):
        ws.cell(row=i, column=1, value=rec.priority).font = _BODY
        ws.cell(row=i, column=2, value=rec.category).font = _BODY
        ws.cell(row=i, column=3, value=rec.title).font = _BODY
        ws.cell(row=i, column=4, value=rec.rationale).font = _BODY
        ws.cell(row=i, column=5, value=rec.recommended_action).font = _BODY
    _autosize(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
