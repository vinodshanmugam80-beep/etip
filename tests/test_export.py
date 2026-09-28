"""Integration test for the board-pack Excel export."""

from __future__ import annotations

import io

import openpyxl
from fastapi.testclient import TestClient

EXP = "/api/v1/exports/board-pack.xlsx"
PROJECTS = "/api/v1/projects"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_board_pack_xlsx(client: TestClient, admin_headers: dict[str, str]) -> None:
    client.post(
        PROJECTS,
        headers=admin_headers,
        json={"name": "Atlas", "code": "ATLAS", "budget": "100000.00"},
    )
    r = client.get(EXP, headers=admin_headers)
    assert r.status_code == 200
    assert XLSX in r.headers["content-type"]
    assert r.headers["content-disposition"].endswith('.xlsx"')

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Projects", "KPIs", "Recommendations"]
    assert wb["Summary"]["A1"].value == "ETIP — Executive Board Pack"
    # The project we created appears on the Projects sheet.
    codes = [
        wb["Projects"].cell(row=i, column=1).value for i in range(2, wb["Projects"].max_row + 1)
    ]
    assert "ATLAS" in codes


def test_board_pack_unauthenticated(client: TestClient) -> None:
    assert client.get(EXP).status_code == 401


def test_board_pack_pdf(client: TestClient, admin_headers: dict[str, str]) -> None:
    client.post(
        PROJECTS,
        headers=admin_headers,
        json={"name": "Atlas", "code": "ATLAS", "budget": "100000.00"},
    )
    r = client.get("/api/v1/exports/board-pack.pdf", headers=admin_headers)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:5] == b"%PDF-"
    assert r.headers["content-disposition"].endswith('.pdf"')
